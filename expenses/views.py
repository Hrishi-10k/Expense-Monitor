import re
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from xml.sax.saxutils import escape

from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.core.validators import validate_email
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import Budget, Category, Expense

MAX_AMOUNT = Decimal('99999999.99')   # fits DecimalField(max_digits=10, decimal_places=2)
ZERO = Decimal('0.00')


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------

def parse_amount(raw):
    """Return a positive Decimal with 2 places, or None if the input is invalid."""
    try:
        value = Decimal(str(raw).strip())
        if not value.is_finite():
            return None
        value = value.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError, TypeError):
        return None
    if value <= 0 or value > MAX_AMOUNT:
        return None
    return value


def parse_iso_date(raw):
    try:
        return parse_date((raw or '').strip())
    except ValueError:
        return None


def to_int(raw):
    raw = (raw or '').strip()
    return int(raw) if raw.isdigit() else None


def clean_name(raw):
    return re.sub(r'\s+', ' ', (raw or '')).strip()


def month_spending(user, year, month):
    """{category_id: total spent} for one calendar month."""
    rows = (
        Expense.objects.filter(
            user=user, expense_date__year=year, expense_date__month=month
        )
        .values('category')
        .annotate(total=Sum('amount'))
    )
    return {row['category']: row['total'] for row in rows}


def read_expense_form(request):
    """Validate the expense form. Returns (data, error_message)."""
    category = Category.objects.filter(
        user=request.user, id=to_int(request.POST.get('category'))
    ).first()
    if category is None:
        return None, "Please choose a valid category."

    amount = parse_amount(request.POST.get('amount'))
    if amount is None:
        return None, "Enter an amount greater than zero (max 99,999,999.99)."

    description = (request.POST.get('description') or '').strip()
    if not description:
        return None, "Description is required."

    expense_date = parse_iso_date(request.POST.get('expense_date'))
    if expense_date is None:
        return None, "Enter a valid date."

    return {
        'category': category,
        'amount': amount,
        'description': description,
        'expense_date': expense_date,
    }, None


def read_budget_form(request, exclude_id=None):
    """Validate the budget form. Returns (data, error_message)."""
    category = Category.objects.filter(
        user=request.user, id=to_int(request.POST.get('category'))
    ).first()
    if category is None:
        return None, "Please choose a valid category."

    limit = parse_amount(request.POST.get('monthly_limit'))
    if limit is None:
        return None, "Enter a monthly budget greater than zero (max 99,999,999.99)."

    duplicate = Budget.objects.filter(user=request.user, category=category)
    if exclude_id is not None:
        duplicate = duplicate.exclude(id=exclude_id)
    if duplicate.exists():
        return None, "Budget already exists for this category!"

    return {'category': category, 'monthly_limit': limit}, None


def category_name_error(request, name, exclude_id=None):
    if not name:
        return "Category name is required."
    if len(name) > 100:
        return "Category name must be 100 characters or fewer."
    duplicate = Category.objects.filter(user=request.user, name__iexact=name)
    if exclude_id is not None:
        duplicate = duplicate.exclude(id=exclude_id)
    if duplicate.exists():
        return "Category already exists!"
    return None


# --------------------------------------------------------------------------
# Dashboard
# --------------------------------------------------------------------------

@login_required
def dashboard(request):
    user = request.user
    today = timezone.localdate()

    expenses = Expense.objects.filter(user=user)

    by_category = list(
        expenses.values('category__name')
        .annotate(total=Sum('amount'))
        .order_by('category__name')
    )
    chart_labels = [row['category__name'] for row in by_category]
    chart_data = [float(row['total']) for row in by_category]

    total_amount = expenses.aggregate(total=Sum('amount'))['total'] or ZERO
    monthly_expense = (
        expenses.filter(expense_date__year=today.year, expense_date__month=today.month)
        .aggregate(total=Sum('amount'))['total']
        or ZERO
    )
    top_category = (
        expenses.values('category__name')
        .annotate(total_spent=Sum('amount'))
        .order_by('-total_spent')
        .first()
    )
    total_budget_amount = (
        Budget.objects.filter(user=user).aggregate(total=Sum('monthly_limit'))['total']
        or ZERO
    )

    # Budgets are MONTHLY limits, so compare them with THIS MONTH's spending.
    if total_budget_amount == 0:
        budget_status = "No Budget Set"
    elif monthly_expense > total_budget_amount:
        budget_status = "Over Budget"
    else:
        budget_status = "Within Budget"

    context = {
        'total_categories': Category.objects.filter(user=user).count(),
        'total_expenses': expenses.count(),
        'total_budgets': Budget.objects.filter(user=user).count(),
        'recent_expenses': expenses.select_related('category')[:5],
        'total_amount': total_amount,
        'monthly_expense': monthly_expense,
        'total_budget_amount': total_budget_amount,
        'budget_status': budget_status,
        'top_category': top_category,
        'chart_labels': chart_labels,
        'chart_data': chart_data,
    }
    return render(request, 'expenses/dashboard.html', context)


# --------------------------------------------------------------------------
# Expenses
# --------------------------------------------------------------------------

@login_required
def add_expense(request):
    if request.method == 'POST':
        data, error = read_expense_form(request)
        if error:
            messages.error(request, error)
            return redirect('add_expense')

        Expense.objects.create(user=request.user, **data)
        messages.success(request, "Expense added successfully!")
        return redirect('dashboard')

    categories = Category.objects.filter(user=request.user)
    return render(request, 'expenses/add_expense.html', {'categories': categories})


@login_required
def view_expenses(request):
    expenses = Expense.objects.filter(user=request.user).select_related('category')

    search = (request.GET.get('search') or '').strip()
    category_id = to_int(request.GET.get('category'))
    date_raw = (request.GET.get('date') or '').strip()
    filter_date = parse_iso_date(date_raw)
    sort = request.GET.get('sort') or ''

    if search:
        expenses = expenses.filter(description__icontains=search)
    if category_id is not None:
        expenses = expenses.filter(category_id=category_id)
    if filter_date:
        expenses = expenses.filter(expense_date=filter_date)

    if sort == 'low':
        expenses = expenses.order_by('amount', '-id')
    elif sort == 'high':
        expenses = expenses.order_by('-amount', '-id')

    page_obj = Paginator(expenses, 20).get_page(request.GET.get('page'))

    # Keep the filters when moving between pages.
    params = request.GET.copy()
    params.pop('page', None)

    context = {
        'page_obj': page_obj,
        'categories': Category.objects.filter(user=request.user),
        'filters': {
            'search': search,
            'category': category_id,
            'date': date_raw if filter_date else '',
            'sort': sort,
        },
        'querystring': params.urlencode(),
    }
    return render(request, 'expenses/view_expenses.html', context)


@login_required
def update_expense(request, id):
    expense = get_object_or_404(Expense, id=id, user=request.user)

    if request.method == 'POST':
        data, error = read_expense_form(request)
        if error:
            messages.error(request, error)
            return redirect('update_expense', id=expense.id)

        for field, value in data.items():
            setattr(expense, field, value)
        expense.save()

        messages.success(request, "Expense updated successfully!")
        return redirect('view_expenses')

    context = {
        'expense': expense,
        'categories': Category.objects.filter(user=request.user),
    }
    return render(request, 'expenses/update_expense.html', context)


@login_required
@require_POST
def delete_expense(request, id):
    expense = get_object_or_404(Expense, id=id, user=request.user)
    expense.delete()
    messages.success(request, "Expense deleted successfully!")
    return redirect('view_expenses')


# --------------------------------------------------------------------------
# Authentication
# --------------------------------------------------------------------------

def register(request):
    if request.user.is_authenticated:
        return redirect('dashboard')

    form_data = {}

    if request.method == 'POST':
        username = (request.POST.get('username') or '').strip()
        email = (request.POST.get('email') or '').strip().lower()
        password = request.POST.get('password') or ''
        form_data = {'username': username, 'email': email}

        errors = []

        if not username:
            errors.append("Username is required.")
        elif len(username) > 150:
            errors.append("Username must be 150 characters or fewer.")
        elif User.objects.filter(username__iexact=username).exists():
            errors.append("Username already exists!")

        try:
            validate_email(email)
        except ValidationError:
            errors.append("Enter a valid email address.")
        else:
            if User.objects.filter(email__iexact=email).exists():
                errors.append("Email already exists!")

        if not errors:
            try:
                validate_password(password, user=User(username=username, email=email))
            except ValidationError as exc:
                errors.extend(exc.messages)

        if errors:
            for error in errors:
                messages.error(request, error)
            return render(request, 'expenses/register.html', {'form_data': form_data})

        User.objects.create_user(username=username, email=email, password=password)
        messages.success(request, "Registration successful! Please login.")
        return redirect('login')

    return render(request, 'expenses/register.html', {'form_data': form_data})


def login_user(request):
    if request.user.is_authenticated:
        return redirect('dashboard')

    next_url = request.POST.get('next') or request.GET.get('next') or ''

    if request.method == 'POST':
        user = authenticate(
            request,
            username=(request.POST.get('username') or '').strip(),
            password=request.POST.get('password') or '',
        )

        if user is not None:
            login(request, user)
            if next_url and url_has_allowed_host_and_scheme(
                next_url,
                allowed_hosts={request.get_host()},
                require_https=request.is_secure(),
            ):
                return redirect(next_url)
            return redirect('dashboard')

        return render(
            request,
            'expenses/login.html',
            {'error': 'Invalid Username or Password', 'next': next_url},
        )

    return render(request, 'expenses/login.html', {'next': next_url})


@require_POST
def logout_user(request):
    logout(request)
    return redirect('login')


# --------------------------------------------------------------------------
# Budgets
# --------------------------------------------------------------------------

@login_required
def add_budget(request):
    if request.method == 'POST':
        data, error = read_budget_form(request)
        if error:
            messages.error(request, error)
            return redirect('add_budget')

        Budget.objects.create(user=request.user, **data)
        messages.success(request, "Budget added successfully!")
        return redirect('dashboard')

    categories = Category.objects.filter(user=request.user)
    return render(request, 'expenses/add_budget.html', {'categories': categories})


@login_required
def view_budgets(request):
    budgets = Budget.objects.filter(user=request.user).select_related('category')
    return render(request, 'expenses/view_budgets.html', {'budgets': budgets})


@login_required
def edit_budget(request, id):
    budget = get_object_or_404(Budget, id=id, user=request.user)

    if request.method == 'POST':
        data, error = read_budget_form(request, exclude_id=budget.id)
        if error:
            messages.error(request, error)
            return redirect('edit_budget', id=budget.id)

        budget.category = data['category']
        budget.monthly_limit = data['monthly_limit']
        budget.save()

        messages.success(request, "Budget updated successfully!")
        return redirect('view_budgets')

    context = {
        'budget': budget,
        'categories': Category.objects.filter(user=request.user),
    }
    return render(request, 'expenses/edit_budget.html', context)


@login_required
@require_POST
def delete_budget(request, id):
    budget = get_object_or_404(Budget, id=id, user=request.user)
    budget.delete()
    messages.success(request, "Budget deleted successfully!")
    return redirect('view_budgets')


@login_required
def budget_report(request):
    """This month's spending against each monthly budget."""
    today = timezone.localdate()
    spending = month_spending(request.user, today.year, today.month)

    report = []
    for budget in Budget.objects.filter(user=request.user).select_related('category'):
        spent = spending.get(budget.category_id, ZERO)
        percent = int(spent / budget.monthly_limit * 100) if budget.monthly_limit else 0
        report.append({
            'category': budget.category.name,
            'budget': budget.monthly_limit,
            'spent': spent,
            'remaining': budget.monthly_limit - spent,
            'percent': percent,
            'bar_percent': min(percent, 100),
        })

    return render(request, 'expenses/budget_report.html', {
        'report': report,
        'month_label': today.strftime('%B %Y'),
    })


# --------------------------------------------------------------------------
# Categories
# --------------------------------------------------------------------------

@login_required
def view_categories(request):
    categories = Category.objects.filter(user=request.user)
    return render(request, 'expenses/view_categories.html', {'categories': categories})


@login_required
def add_category(request):
    if request.method == 'POST':
        name = clean_name(request.POST.get('name'))
        error = category_name_error(request, name)
        if error:
            messages.error(request, error)
            return redirect('add_category')

        Category.objects.create(user=request.user, name=name)
        messages.success(request, "Category added successfully!")
        return redirect('view_categories')

    return render(request, 'expenses/add_category.html')


@login_required
def edit_category(request, id):
    category = get_object_or_404(Category, id=id, user=request.user)

    if request.method == 'POST':
        name = clean_name(request.POST.get('name'))
        error = category_name_error(request, name, exclude_id=category.id)
        if error:
            messages.error(request, error)
            return redirect('edit_category', id=category.id)

        category.name = name
        category.save()
        messages.success(request, "Category updated successfully!")
        return redirect('view_categories')

    return render(request, 'expenses/edit_category.html', {'category': category})


@login_required
@require_POST
def delete_category(request, id):
    category = get_object_or_404(Category, id=id, user=request.user)

    if Expense.objects.filter(category=category).exists():
        messages.error(request, "Cannot delete category because it has expenses.")
        return redirect('view_categories')

    category.delete()
    messages.success(request, "Category deleted successfully!")
    return redirect('view_categories')


# --------------------------------------------------------------------------
# Exports
# --------------------------------------------------------------------------

def _export_rows(user):
    return Expense.objects.filter(user=user).select_related('category').order_by(
        '-expense_date', '-id'
    )


@login_required
def export_pdf(request):
    expenses = list(_export_rows(request.user))
    total = sum((e.amount for e in expenses), ZERO)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="Expense_Report.pdf"'

    doc = SimpleDocTemplate(
        response,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title="Expense Report",
    )
    styles = getSampleStyleSheet()
    cell = styles['BodyText']
    cell.fontSize = 9
    cell.leading = 11

    # Paragraph() reads markup, so user text must be escaped.
    rows = [['Category', 'Amount (INR)', 'Description', 'Date']]
    for e in expenses:
        rows.append([
            Paragraph(escape(e.category.name), cell),
            f"{e.amount:,.2f}",
            Paragraph(escape(e.description).replace('\n', '<br/>'), cell),
            e.expense_date.strftime('%d-%m-%Y'),
        ])
    rows.append(['Total', f"{total:,.2f}", '', ''])

    table = Table(rows, colWidths=[40 * mm, 30 * mm, 80 * mm, 30 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#212529')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#f2f2f2')]),
    ]))

    doc.build([
        Paragraph("Expense Report", styles['Title']),
        Paragraph(f"Generated on {timezone.localdate():%d-%m-%Y}", styles['Normal']),
        Spacer(1, 6 * mm),
        table,
    ])
    return response


@login_required
def export_excel(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Expenses"

    sheet.append(["Category", "Amount", "Description", "Date"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    total = ZERO
    for e in _export_rows(request.user):
        total += e.amount
        sheet.append([e.category.name, float(e.amount), e.description, e.expense_date])
        row = sheet.max_row
        sheet.cell(row=row, column=2).number_format = '#,##0.00'
        sheet.cell(row=row, column=4).number_format = 'DD-MM-YYYY'
        # Text that starts with = + - @ would be run as a formula in Excel
        # (formula injection). Force every text cell to stay plain text.
        for col in (1, 3):
            text_cell = sheet.cell(row=row, column=col)
            if text_cell.data_type == 'f':
                text_cell.data_type = 's'

    sheet.append(["Total", float(total)])
    total_row = sheet.max_row
    sheet.cell(row=total_row, column=1).font = Font(bold=True)
    sheet.cell(row=total_row, column=2).font = Font(bold=True)
    sheet.cell(row=total_row, column=2).number_format = '#,##0.00'

    for letter, width in zip('ABCD', (22, 14, 50, 14)):
        sheet.column_dimensions[letter].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = 'attachment; filename="Expense_Report.xlsx"'
    workbook.save(response)
    return response
