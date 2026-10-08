import io
from datetime import date
from decimal import Decimal

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import Budget, Category, Expense


class BaseCase(TestCase):
    def setUp(self):
        self.alice = User.objects.create_user('alice', 'alice@example.com', 'S3cure-pass!')
        self.bob = User.objects.create_user('bob', 'bob@example.com', 'S3cure-pass!')
        self.food = Category.objects.create(user=self.alice, name='Food')
        self.client.login(username='alice', password='S3cure-pass!')

    def make_expense(self, **kw):
        data = dict(user=self.alice, category=self.food, amount='10.00',
                    description='lunch', expense_date=timezone.localdate())
        data.update(kw)
        return Expense.objects.create(**data)


class AuthTests(BaseCase):
    def test_pages_require_login(self):
        self.client.logout()
        for name in ('dashboard', 'view_expenses', 'budget_report', 'export_pdf', 'export_excel'):
            resp = self.client.get(reverse(name))
            self.assertEqual(resp.status_code, 302, name)
            self.assertIn(reverse('login'), resp['Location'])

    def test_login_redirects_to_next_but_not_to_other_sites(self):
        self.client.logout()
        resp = self.client.post(reverse('login'), {
            'username': 'alice', 'password': 'S3cure-pass!', 'next': reverse('view_budgets')})
        self.assertRedirects(resp, reverse('view_budgets'), fetch_redirect_response=False)
        self.client.logout()
        resp = self.client.post(reverse('login'), {
            'username': 'alice', 'password': 'S3cure-pass!', 'next': 'https://evil.example/'})
        self.assertRedirects(resp, reverse('dashboard'), fetch_redirect_response=False)

    def test_register_rejects_weak_password_and_duplicates(self):
        self.client.logout()
        url = reverse('register')
        resp = self.client.post(url, {'username': 'new', 'email': 'n@example.com', 'password': '123'})
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(User.objects.filter(username='new').exists())
        resp = self.client.post(url, {'username': 'ALICE', 'email': 'x@example.com', 'password': 'S3cure-pass!'})
        self.assertFalse(User.objects.filter(email='x@example.com').exists())
        resp = self.client.post(url, {'username': 'new', 'email': 'ALICE@example.com', 'password': 'S3cure-pass!'})
        self.assertFalse(User.objects.filter(username='new').exists())
        resp = self.client.post(url, {'username': 'new', 'email': 'n@example.com', 'password': 'S3cure-pass!'})
        self.assertRedirects(resp, reverse('login'), fetch_redirect_response=False)

    def test_logout_requires_post(self):
        self.assertEqual(self.client.get(reverse('logout')).status_code, 405)
        self.assertRedirects(self.client.post(reverse('logout')), reverse('login'),
                             fetch_redirect_response=False)


class IsolationTests(BaseCase):
    def test_users_cannot_see_or_touch_each_others_data(self):
        exp = self.make_expense()
        self.client.logout()
        self.client.login(username='bob', password='S3cure-pass!')

        self.assertNotContains(self.client.get(reverse('view_expenses')), 'lunch')
        self.assertEqual(self.client.get(reverse('update_expense', args=[exp.id])).status_code, 404)
        self.assertEqual(self.client.post(reverse('delete_expense', args=[exp.id])).status_code, 404)
        self.assertEqual(self.client.post(reverse('delete_category', args=[self.food.id])).status_code, 404)
        self.assertTrue(Expense.objects.filter(id=exp.id).exists())

        # Bob cannot file an expense under Alice's category.
        resp = self.client.post(reverse('add_expense'), {
            'category': self.food.id, 'amount': '5', 'description': 'x',
            'expense_date': '2026-01-01'})
        self.assertEqual(Expense.objects.filter(user=self.bob).count(), 0)

    def test_same_category_name_allowed_for_different_users(self):
        self.client.logout()
        self.client.login(username='bob', password='S3cure-pass!')
        self.client.post(reverse('add_category'), {'name': 'Food'})
        self.assertTrue(Category.objects.filter(user=self.bob, name='Food').exists())


class ExpenseTests(BaseCase):
    def post_expense(self, **kw):
        data = {'category': self.food.id, 'amount': '12.50', 'description': 'tea',
                'expense_date': '2026-03-04'}
        data.update(kw)
        return self.client.post(reverse('add_expense'), data)

    def test_add_valid_expense_uses_decimal(self):
        self.post_expense()
        e = Expense.objects.get()
        self.assertEqual(e.amount, Decimal('12.50'))
        self.assertEqual(e.user, self.alice)

    def test_invalid_input_is_rejected_without_crashing(self):
        for bad in ({'amount': '0'}, {'amount': '-5'}, {'amount': 'abc'}, {'amount': ''},
                    {'amount': 'NaN'}, {'amount': '1e400'}, {'amount': '999999999'},
                    {'description': '   '}, {'expense_date': 'nope'},
                    {'expense_date': '2026-13-45'}, {'category': 'zzz'}, {'category': 99999}):
            resp = self.post_expense(**bad)
            self.assertEqual(resp.status_code, 302, bad)
        self.assertEqual(Expense.objects.count(), 0)

    def test_update_validates_too(self):
        e = self.make_expense()
        url = reverse('update_expense', args=[e.id])
        self.client.post(url, {'category': self.food.id, 'amount': '-1',
                               'description': 'x', 'expense_date': '2026-01-01'})
        e.refresh_from_db()
        self.assertEqual(e.amount, Decimal('10.00'))
        self.client.post(url, {'category': self.food.id, 'amount': '20',
                               'description': 'dinner', 'expense_date': '2026-01-02'})
        e.refresh_from_db()
        self.assertEqual((e.amount, e.description), (Decimal('20.00'), 'dinner'))

    def test_delete_needs_post(self):
        e = self.make_expense()
        self.assertEqual(self.client.get(reverse('delete_expense', args=[e.id])).status_code, 405)
        self.assertTrue(Expense.objects.filter(id=e.id).exists())
        self.client.post(reverse('delete_expense', args=[e.id]))
        self.assertFalse(Expense.objects.filter(id=e.id).exists())

    def test_filters_survive_garbage_input(self):
        self.make_expense()
        resp = self.client.get(reverse('view_expenses'), {'date': 'garbage', 'category': 'x', 'page': 'zz'})
        self.assertEqual(resp.status_code, 200)
        resp = self.client.get(reverse('view_expenses'), {'search': 'lun', 'sort': 'high'})
        self.assertContains(resp, 'lunch')

    def test_missing_objects_give_404_not_500(self):
        for name in ('update_expense', 'edit_budget', 'edit_category'):
            self.assertEqual(self.client.get(reverse(name, args=[9999])).status_code, 404)


class CategoryTests(BaseCase):
    def test_duplicate_names_blocked_case_insensitively_on_add_and_edit(self):
        self.client.post(reverse('add_category'), {'name': ' food '})
        self.assertEqual(Category.objects.filter(user=self.alice).count(), 1)
        other = Category.objects.create(user=self.alice, name='Travel')
        self.client.post(reverse('edit_category', args=[other.id]), {'name': 'FOOD'})
        other.refresh_from_db()
        self.assertEqual(other.name, 'Travel')

    def test_cannot_delete_category_with_expenses(self):
        self.make_expense()
        self.client.post(reverse('delete_category', args=[self.food.id]))
        self.assertTrue(Category.objects.filter(id=self.food.id).exists())


class BudgetTests(BaseCase):
    def test_duplicate_budget_blocked_on_add_and_edit(self):
        travel = Category.objects.create(user=self.alice, name='Travel')
        self.client.post(reverse('add_budget'), {'category': self.food.id, 'monthly_limit': '100'})
        self.client.post(reverse('add_budget'), {'category': self.food.id, 'monthly_limit': '200'})
        self.assertEqual(Budget.objects.count(), 1)
        self.client.post(reverse('add_budget'), {'category': travel.id, 'monthly_limit': '50'})
        b = Budget.objects.get(category=travel)
        self.client.post(reverse('edit_budget', args=[b.id]),
                         {'category': self.food.id, 'monthly_limit': '50'})
        b.refresh_from_db()
        self.assertEqual(b.category, travel)

    def test_report_only_counts_the_current_month(self):
        Budget.objects.create(user=self.alice, category=self.food, monthly_limit='100')
        today = timezone.localdate()
        self.make_expense(amount='30.00', expense_date=today)
        last_year = today.replace(year=today.year - 1)
        self.make_expense(amount='500.00', expense_date=last_year)
        report = self.client.get(reverse('budget_report')).context['report']
        self.assertEqual(report[0]['spent'], Decimal('30.00'))
        self.assertEqual(report[0]['remaining'], Decimal('70.00'))

    def test_dashboard_status_uses_monthly_spending(self):
        Budget.objects.create(user=self.alice, category=self.food, monthly_limit='100')
        today = timezone.localdate()
        self.make_expense(amount='500.00', expense_date=today.replace(year=today.year - 1))
        ctx = self.client.get(reverse('dashboard')).context
        self.assertEqual(ctx['budget_status'], 'Within Budget')
        self.make_expense(amount='150.00', expense_date=today)
        ctx = self.client.get(reverse('dashboard')).context
        self.assertEqual(ctx['budget_status'], 'Over Budget')


class DashboardAndExportTests(BaseCase):
    def test_dashboard_chart_data_is_safe_json(self):
        evil = Category.objects.create(user=self.alice, name="</script><script>alert('x')</script>")
        self.make_expense(category=evil)
        resp = self.client.get(reverse('dashboard'))
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, "</script><script>alert")

    def test_dashboard_shows_only_five_recent(self):
        for i in range(8):
            self.make_expense(description=f'item{i}')
        self.assertEqual(len(self.client.get(reverse('dashboard')).context['recent_expenses']), 5)

    def test_excel_export_neutralises_formulas(self):
        self.make_expense(description='=HYPERLINK("http://evil","x")', amount='12.34')
        resp = self.client.get(reverse('export_excel'))
        ws = load_workbook(io.BytesIO(resp.content)).active
        self.assertEqual(ws['C2'].data_type, 's')
        self.assertEqual(ws['B2'].value, 12.34)
        self.assertEqual(ws['B3'].value, 12.34)   # total row

    def test_exports_cover_only_own_data_and_many_rows(self):
        for i in range(120):
            self.make_expense(description='very long text ' * 20 + str(i))
        Expense.objects.create(user=self.bob, category=Category.objects.create(user=self.bob, name='B'),
                               amount='1', description='bobs secret', expense_date=date.today())
        pdf = self.client.get(reverse('export_pdf'))
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b'%PDF'))
        ws = load_workbook(io.BytesIO(self.client.get(reverse('export_excel')).content)).active
        self.assertEqual(ws.max_row, 1 + 120 + 1)

    def test_pdf_handles_markup_characters(self):
        self.make_expense(description='<b>bold & <unclosed', amount='1')
        self.assertEqual(self.client.get(reverse('export_pdf')).status_code, 200)


class PageRenderTests(BaseCase):
    def test_every_page_renders(self):
        exp = self.make_expense()
        bud = Budget.objects.create(user=self.alice, category=self.food, monthly_limit='100')
        pages = [
            reverse('dashboard'), reverse('add_expense'), reverse('view_expenses'),
            reverse('update_expense', args=[exp.id]), reverse('add_budget'),
            reverse('view_budgets'), reverse('edit_budget', args=[bud.id]),
            reverse('budget_report'), reverse('view_categories'), reverse('add_category'),
            reverse('edit_category', args=[self.food.id]),
        ]
        for url in pages:
            resp = self.client.get(url)
            self.assertEqual(resp.status_code, 200, url)
            self.assertContains(resp, 'Logout')          # sidebar for signed-in users
        self.client.logout()
        for url in (reverse('login'), reverse('register')):
            resp = self.client.get(url)
            self.assertEqual(resp.status_code, 200, url)
            self.assertNotContains(resp, 'Export PDF')   # no sidebar when signed out

    def test_error_messages_use_bootstrap_danger_class(self):
        self.client.post(reverse('add_category'), {'name': ''})
        resp = self.client.get(reverse('view_categories'))
        self.assertContains(resp, 'alert-danger')
        self.assertNotContains(resp, 'alert-error')
